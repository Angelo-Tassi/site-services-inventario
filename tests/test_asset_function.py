"""La colonna Asset Function: Standard o PC Refresh.

Quello che e' stato chiesto, parola per parola:
- due valori soli, Standard e PC Refresh;
- nell'ordine, subito a destra di Tipo e prima di Stanza;
- esce da tutte le esportazioni;
- non e' obbligatoria in nessuna importazione.

E quello che ne segue: un valore scritto male ma riconoscibile si raddrizza,
uno che non e' nessuno dei due resta vuoto e l'importazione lo dice; un
dispositivo che non ce l'ha resta vuoto, non si inventa; gli iPhone non ce
l'hanno.
"""
import os, sys, tempfile, time
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import fixture
from openpyxl import Workbook, load_workbook
from tkinter import messagebox
from inventario import excel_io
from inventario.store import (FUNZIONI, HEADERS, InventoryError, InventoryStore,
                              funzione_canonica, new_item, rows_from_workbook)
from inventario.ui import App, ImportDialog

BAU, KIOSK, DR = fixture.BAU, fixture.KIOSK, fixture.DR
STANZE = [BAU, KIOSK, DR]

# ---- due valori soli
assert FUNZIONI == ["Standard", "PC Refresh"], FUNZIONI
assert HEADERS["funzione"] == "Asset Function"
for scritto, atteso in (("standard", "Standard"), ("STANDARD", "Standard"),
                        ("pc refresh", "PC Refresh"), ("PC-Refresh", "PC Refresh"),
                        ("pcrefresh", "PC Refresh"), ("Refresh", "PC Refresh"),
                        ("Ricondizionato", ""), ("", "")):
    assert funzione_canonica(scritto) == atteso, (scritto, funzione_canonica(scritto))

# ---- gli iPhone non ce l'hanno
assert new_item("X", "Iphone", imei="356938035643809", funzione="Standard")["funzione"] == ""

# ======================= importare: mai obbligatoria =======================
def foglio(intestazioni, righe):
    wb = Workbook(); ws = wb.active; ws.title = "Inventario"
    ws.append(intestazioni)
    for r in righe:
        ws.append(r)
    p = os.path.join(tempfile.mkdtemp(), "f.xlsx"); wb.save(p); wb.close()
    return p

# ---- un foglio senza la colonna si importa lo stesso, e la lascia vuota
senza = foglio(["Asset Tag", "Tipo", "Stanza"],
               [["IT-9101", "Laptop", DR], ["IT-9102", "Tablet", DR]])
items, esito = rows_from_workbook(senza, STANZE)
assert len(items) == 2 and all(i["funzione"] == "" for i in items), items
assert esito["funzione_sconosciuta"] == []
store = InventoryStore(fixture.build(), iphone_room=BAU)
store.stanze = STANZE
store.load()
assert store.import_items(items, "merge")["aggiunti"] == 2
store.load()
assert [i["funzione"] for i in store.items if i["asset_tag"] in ("IT-9101", "IT-9102")] == ["", ""]

# ---- con la colonna: i valori si raddrizzano, e uno sbagliato si segnala
con = foglio(["Asset Tag", "Tipo", "Asset Function", "Stanza"],
             [["IT-9201", "Laptop", "pc refresh", DR],
              ["IT-9202", "Laptop", "Standard", DR],
              ["IT-9203", "Laptop", "", DR],
              ["IT-9204", "Laptop", "Ricondizionato", DR]])
items, esito = rows_from_workbook(con, STANZE)
funz = dict((i["asset_tag"], i["funzione"]) for i in items)
assert funz == {"IT-9201": "PC Refresh", "IT-9202": "Standard",
                "IT-9203": "", "IT-9204": ""}, funz
assert esito["funzione_sconosciuta"] == ["IT-9204: Ricondizionato"], esito
avvisi = ImportDialog._avvertenze(esito)
assert any("Ricondizionato" in a and "non e' obbligatorio" in a for a in avvisi), avvisi
# ...e la riga con il valore sbagliato entra comunque: il campo non e' obbligatorio
assert store.import_items(items, "merge")["aggiunti"] == 4

# ---- anche il nome della colonna si riconosce in piu' forme
for nome in ("asset function", "Funzione", "ASSET FUNCTION"):
    p = foglio(["Asset Tag", "Tipo", nome, "Stanza"], [["IT-9301", "Laptop", "Standard", DR]])
    items, _e = rows_from_workbook(p, STANZE)
    assert items[0]["funzione"] == "Standard", (nome, items)

# ---- e il vecchio file dati senza la colonna si apre, con la funzione vuota
vecchio = InventoryStore(fixture.build(), iphone_room=BAU)
vecchio.load()
assert all(i.get("funzione", "") == "" for i in vecchio.items)

# ======================= esportare: sempre, fra Tipo e Stanza =======================
store.load()
fuori = tempfile.mkdtemp()

def intestazioni(percorso, foglio_n=0):
    wb = load_workbook(percorso)
    try:
        for riga in wb.worksheets[foglio_n].iter_rows(values_only=True):
            if riga and riga[0] and str(riga[0]).strip() == "Asset Tag":
                return [c for c in riga if c]
    finally:
        wb.close()

def nel_posto_giusto(colonne, dove):
    assert "Asset Function" in colonne, (dove, colonne)
    i = colonne.index("Asset Function")
    assert colonne[i - 1] == "Tipo", (dove, colonne)
    if "Stanza" in colonne:
        # prima della Stanza: subito prima dove le due si toccano, e comunque
        # prima nella stampa, che fra le due ha anche modello e seriale
        assert i < colonne.index("Stanza"), (dove, colonne)

tutto = excel_io.export(store.items, os.path.join(fuori, "tutto.xlsx"), rooms=STANZE)
nel_posto_giusto(intestazioni(tutto), "un unico elenco")
per_stanza = excel_io.export(store.items, os.path.join(fuori, "stanze.xlsx"),
                             group_by_room=True, rooms=STANZE)
nel_posto_giusto(intestazioni(per_stanza), "un foglio per stanza")
for f in excel_io.export_per_stanza(store.items, fuori, rooms=STANZE):
    nel_posto_giusto(intestazioni(f), "un file per stanza")
nel_posto_giusto(intestazioni(excel_io.build_print_file(store.items)), "la stampa")
inglese = excel_io.export(store.items, os.path.join(fuori, "en.xlsx"), rooms=STANZE, lingua="en")
colonne_en = intestazioni(inglese)
assert colonne_en[colonne_en.index("Asset Function") - 1] == "Type", colonne_en
# esce anche quando nessuno ce l'ha: e' una colonna dell'inventario, non un avanzo
nessuno = [dict(i, funzione="") for i in store.items]
nel_posto_giusto(intestazioni(excel_io.export(nessuno, os.path.join(fuori, "vuota.xlsx"),
                                              rooms=STANZE)), "tutte vuote")

# ---- il valore esce com'e'
wb = load_workbook(tutto)
righe = list(wb.worksheets[0].iter_rows(values_only=True))
wb.close()
testa_i = [i for i, r in enumerate(righe) if r and r[0] == "Asset Tag"][0]
col = list(righe[testa_i]).index("Asset Function")
valori = dict((r[0], r[col]) for r in righe[testa_i + 1:] if r and r[0])
assert valori.get("IT-9201") == "PC Refresh" and valori.get("IT-9202") == "Standard", valori

# ---- il modello da compilare ce l'ha, con la sua tendina
modello = os.path.join(fuori, "modello.xlsx")
excel_io.build_template(modello, STANZE)
nel_posto_giusto(intestazioni(modello), "il modello")
wb = load_workbook(modello)
tendine = [dv.formula1 for dv in wb["Inventario"].data_validations.dataValidation]
wb.close()
assert '"Standard,PC Refresh"' in tendine, tendine

# ---- ed esportare e reimportare non perde niente
items, _e = rows_from_workbook(tutto, STANZE)
dopo = dict((i["asset_tag"], i["funzione"]) for i in items)
assert dopo["IT-9201"] == "PC Refresh" and dopo["IT-9202"] == "Standard", dopo

# ======================= nel programma =======================
app = App(fixture.build())
app._initial_load()
app.update()
scadenza = time.time() + 0.4
while time.time() < scadenza:
    app.update()
for nome in ("showinfo", "showwarning", "showerror"):
    setattr(messagebox, nome, lambda t, m, **k: None)

# ---- la colonna c'e', fra Tipo e Stanza
app.show_home(); app.update()
campi = app._campi_visibili()
assert campi[campi.index("funzione") - 1] == "tipo" and \
    campi[campi.index("funzione") + 1] == "stanza", campi
# ...anche dentro una stanza, dove la Stanza non c'e'
app.show_room(KIOSK); app.update()
campi = app._campi_visibili()
assert campi[campi.index("funzione") - 1] == "tipo", campi
# ...ma non nel contenitore degli iPhone
app.show_iphones(); app.update()
assert "funzione" not in app._campi_visibili()

# ---- si cambia dall'archivio, anche a un dispositivo in prestito
tag = "IT-0107"                       # in prestito nella fixture
assert app.store.set_funzione(tag, "pc refresh") is True
app.store.load()
assert [i for i in app.store.items if i["asset_tag"] == tag][0]["funzione"] == "PC Refresh"
try:
    app.store.set_funzione(tag, "Ricondizionato")
    raise SystemExit("valore non previsto accettato")
except InventoryError:
    pass

# ---- dalla scheda: parte da Standard per un dispositivo nuovo, tiene la sua
from inventario.ui import ItemDialog
d = ItemDialog(app, app.cfg["rooms"], app.cfg["types"], iphone_room=BAU,
               stati=app.cfg["states"])
assert d.var_funzione.get() == "Standard"
d.destroy()
esistente = dict([i for i in app.store.items if i["asset_tag"] == "IT-0101"][0])
d = ItemDialog(app, app.cfg["rooms"], app.cfg["types"], item=esistente,
               iphone_room=BAU, stati=app.cfg["states"])
assert d.var_funzione.get() == "", "chi non ce l'ha resta vuoto: non si inventa"
d.var_funzione.set("PC Refresh")
d.var_tag.set("IT-0101")
d._ok()
assert d.result["funzione"] == "PC Refresh", d.result

# ---- e la ricerca la trova
app.show_home(); app.update()
app.var_search.set("pc refresh"); app.update()
assert tag in [i["asset_tag"] for i in app.visible], "la ricerca guarda anche la funzione"
app.var_search.set("")

app.destroy()
print("ASSET FUNCTION OK")
