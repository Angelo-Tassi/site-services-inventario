"""Una postazione vecchia avvisa, consulta ed esporta senza riscrivere i dati."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import fixture
from openpyxl import load_workbook
from inventario import __version__, config
from inventario.store import VERSION_SHEET
from inventario.ui import App

percorso = fixture.build()


def imposta_versione(valore):
    wb = load_workbook(percorso)
    wb[VERSION_SHEET]["B1"] = valore
    wb.save(percorso)
    wb.close()


imposta_versione("2.0.0")
prima = open(percorso, "rb").read()
impostazioni = open(config.shared_config_path(percorso), "rb").read()
fixture.popup.clear()
app = App(percorso)
try:
    app._initial_load()
    app.update()
    assert app.store.sola_lettura
    assert "Sola lettura" in app.title()
    assert len(app.tree.get_children()) == 13
    avvisi = [p for p in fixture.popup if p[0] == "showwarning"]
    assert len(avvisi) == 1 and "2.0.0" in avvisi[0][2], avvisi
    assert "aggiorna il programma" in avvisi[0][2]
    app._run(lambda: app.store.set_note("IT-0101", "vietata"))
    assert open(percorso, "rb").read() == prima
    assert open(config.shared_config_path(percorso), "rb").read() == impostazioni
    app.ricostruisci()
    app.update()
    assert "Sola lettura" in app.title()
    assert len([p for p in fixture.popup if p[0] == "showwarning"]) == 1
    # Le copie in locale restano permesse e conservano il marcatore futuro.
    copia = os.path.join(os.path.dirname(percorso), "consultazione.xlsx")
    app.store.copia_in(copia)
    wb = load_workbook(copia)
    assert wb[VERSION_SHEET]["B1"].value == "2.0.0"
    wb.close()
    imposta_versione(__version__)
    app._reload()
    assert not app.store.sola_lettura
    assert "Sola lettura" not in app.title()
finally:
    app.destroy()

# La diagnostica apre l'interfaccia senza un popup bloccante.
fixture.popup.clear()
imposta_versione("2.0.0")
app = App(percorso, avvisi_versione=False)
try:
    app._initial_load()
    assert app.store.sola_lettura and not [p for p in fixture.popup if p[0] == "showwarning"]
finally:
    app.destroy()

print("VERSIONE UI OK")
