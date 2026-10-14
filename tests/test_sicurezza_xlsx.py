"""Colonne future e versioni: nessuna postazione cancella dati che non conosce."""
import os
import sys
import tempfile
from datetime import datetime
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from openpyxl import Workbook, load_workbook
from inventario import __version__, config, store
from inventario.store import InventoryError, InventoryStore, VERSION_SHEET, new_item, ordine_versione


def rifiutata(azione):
    try:
        azione()
    except InventoryError as exc:
        assert "sola lettura" in str(exc), exc
    else:
        raise AssertionError("Scrittura di una versione vecchia consentita")


def versione(percorso, valore):
    wb = load_workbook(percorso)
    if VERSION_SHEET not in wb.sheetnames:
        wb.create_sheet(VERSION_SHEET)
    wb[VERSION_SHEET]["B1"] = valore
    wb[VERSION_SHEET].sheet_state = "hidden"
    wb.save(percorso)
    wb.close()


def contenuto(percorso):
    wb = load_workbook(percorso)
    try:
        return list(wb["Inventario"].values)
    finally:
        wb.close()


with tempfile.TemporaryDirectory() as cartella:
    percorso = os.path.join(cartella, "Inventario.xlsx")
    backup = os.path.join(cartella, "Backup")
    os.mkdir(backup)
    config.backup_dir = lambda: backup
    # File precedente al marcatore di versione, con valori non testuali,
    # formule, intestazioni duplicate e una riga fuori ordine.
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(["Asset Tag", "Tipo", "Stanza", "Asset Function", "Campo futuro", "Costo", "Calcolo", "Campo futuro"])
    data = datetime(2026, 10, 14)
    ws.append(["B", "Laptop", "Sala", "PC Refresh", "valore B", 0, "=1+2", data])
    ws.append(["A", "Laptop", "Sala", "Standard", "valore A", False, "=2+3", None])
    wb.save(percorso)
    wb.close()
    archivio = InventoryStore(percorso)
    archivio.load()
    assert not archivio.sola_lettura
    archivio.set_note("B", "aggiornata")
    archivio.add(new_item("C", "Tablet", stanza="Sala"))
    aggiornato = dict(next(i for i in archivio.items if i["asset_tag"] == "B"), asset_tag="D")
    archivio.update("B", aggiornato)
    righe = contenuto(percorso)
    assert righe[0][-4:] == ("Campo futuro", "Costo", "Calcolo", "Campo futuro")
    per_tag = {r[0]: r for r in righe[1:]}
    assert per_tag["D"][-4:] == ("valore B", 0, "=1+2", data), per_tag
    assert per_tag["A"][-4:] == ("valore A", False, "=2+3", None)
    assert per_tag["C"][-4:] == (None,) * 4
    assert InventoryStore(percorso).load()[0]["funzione"] == "Standard"
    # Anche una postazione che non conoscesse Asset Function la mantiene.
    with patch.object(store, "ALL_FIELDS", [f for f in store.ALL_FIELDS if f != "funzione"]), \
            patch.object(store, "HEADER_ALIASES", {f: v for f, v in store.HEADER_ALIASES.items() if f != "funzione"}):
        InventoryStore(percorso).set_note("D", "da postazione senza colonna")
    assert next(i for i in InventoryStore(percorso).load() if i["asset_tag"] == "D")["funzione"] == "PC Refresh"
    # I dati delle colonne eliminate non ricompaiono su un nuovo inserimento.
    archivio.delete(["D"])
    archivio.add(new_item("D", "Laptop", stanza="Sala"))
    righe = contenuto(percorso)
    nuovo = next(r for r in righe[1:] if r[0] == "D")
    assert all(nuovo[i] is None for i, titolo in enumerate(righe[0])
               if titolo in ("Campo futuro", "Costo", "Calcolo")), nuovo
    # Schema riletto sotto lock: una colonna comparsa dopo load non sparisce.
    wb = load_workbook(percorso)
    ws = wb["Inventario"]
    ws.cell(1, ws.max_column + 1, "Ancora nuova")
    ws.cell(2, ws.max_column, "comparsa dopo load")
    wb.save(percorso)
    wb.close()
    archivio.set_note("A", "riletta")
    assert contenuto(percorso)[1][-1] == "comparsa dopo load"

    wb = load_workbook(percorso)
    assert wb[VERSION_SHEET]["B1"].value == __version__
    assert wb[VERSION_SHEET].sheet_state == "hidden"
    wb.close()
    copia = os.path.join(cartella, "copia.xlsx")
    archivio.copia_in(copia)
    # Il ripristino conserva i campi estesi e registra la postazione che lo
    # esegue, anche se la copia risale a prima del marcatore.
    versione(copia, None)
    prima_copia = contenuto(copia)
    archivio.restore(copia)
    assert contenuto(percorso) == prima_copia
    wb = load_workbook(percorso)
    assert wb[VERSION_SHEET]["B1"].value == __version__
    wb.close()
    # La versione cambia mentre la vecchia postazione e' gia' aperta: una
    # sempre piu' nuova di quella in uso, qualunque sia.
    versione(percorso, "99.0.0-beta.1")
    prima = open(percorso, "rb").read()
    for azione in (lambda: archivio.set_note("A", "vietata"),
                   lambda: archivio.import_items([], "replace"),
                   archivio.reset, lambda: archivio.restore(copia),
                   lambda: archivio.ripristina_da_copia_locale(copia),
                   lambda: config.save_shared_config(percorso, {"rooms": ["Altra"]}),
                   lambda: archivio._write([])):
        rifiutata(azione)
        assert open(percorso, "rb").read() == prima
    altro = InventoryStore(percorso)
    assert altro.load() and altro.sola_lettura
    assert "99.0.0-beta.1" in altro.avviso_versione()
    # Nemmeno una copia piu' nuova puo' essere ripristinata da codice vecchio.
    versione(percorso, __version__)
    versione(copia, "2.0.0")
    prima = open(percorso, "rb").read()
    rifiutata(lambda: archivio.restore(copia))
    rifiutata(lambda: archivio.ripristina_da_copia_locale(copia))
    assert open(percorso, "rb").read() == prima
    for vecchia in (None, "1.0.0-beta.5.8", __version__):
        versione(percorso, vecchia)
        InventoryStore(percorso).set_note("A", "consentita")
    versione(percorso, "versione illeggibile")
    rifiutata(lambda: InventoryStore(percorso).set_note("A", "vietata"))

assert ordine_versione("1.0.0-beta.5.9") < ordine_versione("1.0.0-beta.5.10")
assert ordine_versione("1.0.0-beta.99") < ordine_versione("1.0.0")
assert ordine_versione("1.0.0") < ordine_versione("1.0.1-alpha.1")
assert ordine_versione("1.0.0+build.2") == ordine_versione("1.0.0")
print("SICUREZZA XLSX OK")
